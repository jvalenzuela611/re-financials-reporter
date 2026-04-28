"""
Management Commentary Module — Persistent analyst input per asset.

Allows analysts to add qualitative context (recurring themes, one-off events)
that enriches the AI-generated Notes to Financials. Commentary is stored in the
"Management Commentary" sheet of the per-year Master Excel file.

Each entry is a row with: account/line, period, type (recurring/one_off), text.
Entries accumulate across quarters and serve as the "memory layer" for each asset.
"""

import re
import pandas as pd
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict

from openpyxl import load_workbook


SHEET_NAME = "Management Commentary"

# Column order in the Excel sheet
_COLUMNS = [
    'account_code', 'line_name', 'level', 'period',
    'comment_type', 'text', 'author', 'timestamp',
]


@dataclass
class CommentaryEntry:
    account_code: str      # GL code ("" for L1-level comments)
    line_name: str         # "Income", "Operating Expenses", or L3 account name
    level: str             # "L1" or "L3"
    period: str            # "Q1 2025"
    comment_type: str      # "recurring" or "one_off"
    text: str              # The analyst's commentary
    author: str            # Who wrote it
    timestamp: str         # ISO datetime string


def _master_path(building: str, year: int, master_dir: str) -> Path:
    from src.master_excel import _sanitize_building_name
    safe_name = _sanitize_building_name(building)
    return Path(master_dir) / f"master_{safe_name}_{year}.xlsx"


def _extract_year(period: str) -> int:
    match = re.search(r'(\d{4})', period)
    return int(match.group(1)) if match else datetime.now().year


def load_commentary(
    building: str,
    year: int,
    master_dir: str = "./masters",
) -> List[CommentaryEntry]:
    """Load all management commentary entries from the Master Excel."""
    path = _master_path(building, year, master_dir)
    if not path.exists():
        return []

    try:
        df = pd.read_excel(str(path), sheet_name=SHEET_NAME, dtype=str)
    except (ValueError, KeyError):
        # Sheet doesn't exist yet
        return []

    entries = []
    for _, row in df.iterrows():
        entries.append(CommentaryEntry(
            account_code=str(row.get('account_code', '')).strip(),
            line_name=str(row.get('line_name', '')).strip(),
            level=str(row.get('level', 'L1')).strip(),
            period=str(row.get('period', '')).strip(),
            comment_type=str(row.get('comment_type', 'recurring')).strip(),
            text=str(row.get('text', '')).strip(),
            author=str(row.get('author', '')).strip(),
            timestamp=str(row.get('timestamp', '')).strip(),
        ))
    return entries


def save_commentary(
    building: str,
    year: int,
    entries: List[CommentaryEntry],
    master_dir: str = "./masters",
) -> None:
    """Write management commentary entries to the Master Excel.

    Replaces the entire "Management Commentary" sheet with the provided entries.
    Preserves all other sheets.
    """
    path = _master_path(building, year, master_dir)
    if not path.exists():
        return

    df = pd.DataFrame([asdict(e) for e in entries], columns=_COLUMNS)

    wb = load_workbook(str(path))

    # Remove existing sheet if present
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]

    wb.save(str(path))

    # Write the new sheet using pandas (appends to existing workbook)
    with pd.ExcelWriter(str(path), engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        df.to_excel(writer, sheet_name=SHEET_NAME, index=False)


def add_entry(
    building: str,
    year: int,
    entry: CommentaryEntry,
    master_dir: str = "./masters",
) -> List[CommentaryEntry]:
    """Add a single commentary entry and return the updated list."""
    entries = load_commentary(building, year, master_dir)
    entries.append(entry)
    save_commentary(building, year, entries, master_dir)
    return entries


def delete_entry(
    building: str,
    year: int,
    index: int,
    master_dir: str = "./masters",
) -> List[CommentaryEntry]:
    """Delete a commentary entry by index and return the updated list."""
    entries = load_commentary(building, year, master_dir)
    if 0 <= index < len(entries):
        entries.pop(index)
        save_commentary(building, year, entries, master_dir)
    return entries


def get_commentary_for_prompt(
    building: str,
    year: int,
    master_dir: str = "./masters",
) -> str:
    """Format all commentary entries as structured text for AI prompt injection.

    Returns empty string if no commentary exists.
    """
    entries = load_commentary(building, year, master_dir)
    if not entries:
        return ""

    recurring = [e for e in entries if e.comment_type == 'recurring']
    one_off = [e for e in entries if e.comment_type == 'one_off']

    parts = []

    if recurring:
        parts.append("RECURRING ITEMS (apply to every quarter unless resolved):")
        for e in recurring:
            location = f"[{e.line_name}]"
            if e.account_code:
                location = f"[{e.line_name}, GL {e.account_code}]"
            parts.append(f"  {location} {e.text}")

    if one_off:
        parts.append("")
        parts.append("ONE-OFF EVENTS (specific period only):")
        for e in one_off:
            location = f"[{e.line_name}, {e.period}]"
            if e.account_code:
                location = f"[{e.line_name}, GL {e.account_code}, {e.period}]"
            parts.append(f"  {location} {e.text}")

    return "\n".join(parts)
